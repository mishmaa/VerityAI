import hashlib
import math
import re
import shutil
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import UploadFile
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import DocumentChunk, DocumentStatus, KnowledgeDocument
from app.schemas.documents import Citation

ALLOWED_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "text/markdown": ".md",
}
TOKEN_RE = re.compile(r"[a-zA-Z0-9][a-zA-Z0-9_-]+")
QUERY_STOP_WORDS = {
    "a",
    "an",
    "and",
    "are",
    "as",
    "at",
    "be",
    "by",
    "do",
    "does",
    "for",
    "from",
    "how",
    "i",
    "in",
    "is",
    "it",
    "of",
    "on",
    "or",
    "our",
    "the",
    "their",
    "this",
    "to",
    "was",
    "what",
    "when",
    "where",
    "which",
    "who",
    "why",
    "with",
}


class InvalidDocument(ValueError):
    pass


@dataclass(frozen=True)
class PageText:
    page_number: int | None
    text: str


def safe_filename(name: str | None) -> str:
    value = Path(name or "document").name
    cleaned = re.sub(r"[^\w.() -]", "_", value, flags=re.UNICODE).strip(" .")
    return cleaned[:255] or "document"


def store_upload(upload: UploadFile, organization_id: uuid.UUID) -> tuple[str, int, str]:
    content_type = (upload.content_type or "").lower()
    extension = ALLOWED_TYPES.get(content_type)
    if extension is None:
        raise InvalidDocument("Supported formats are PDF, DOCX, TXT, and Markdown")

    root = Path(settings.upload_dir).resolve()
    tenant_dir = (root / str(organization_id)).resolve()
    if root not in tenant_dir.parents:
        raise InvalidDocument("Invalid storage path")
    tenant_dir.mkdir(parents=True, exist_ok=True)
    destination = tenant_dir / f"{uuid.uuid4()}{extension}"

    digest = hashlib.sha256()
    size = 0
    try:
        with destination.open("xb") as target:
            while block := upload.file.read(1024 * 1024):
                size += len(block)
                if size > settings.max_upload_bytes:
                    raise InvalidDocument(
                        f"File exceeds the {settings.max_upload_bytes // (1024 * 1024)} MB limit"
                    )
                digest.update(block)
                target.write(block)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        upload.file.close()

    if size == 0:
        destination.unlink(missing_ok=True)
        raise InvalidDocument("The uploaded file is empty")
    return str(destination), size, digest.hexdigest()


def extract_pages(path: str, content_type: str) -> list[PageText]:
    if content_type == "application/pdf":
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise InvalidDocument("Password-protected PDFs are not supported") from exc
        return [
            PageText(page_number=index, text=(page.extract_text() or "").strip())
            for index, page in enumerate(reader.pages, start=1)
        ]

    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = DocxDocument(path)
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        )
        return [PageText(page_number=None, text=text)]

    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return [PageText(page_number=1, text=text)]


def chunk_pages(
    pages: list[PageText], target_words: int = 220, overlap: int = 35
) -> list[PageText]:
    chunks: list[PageText] = []
    for page in pages:
        words = page.text.split()
        if not words:
            continue
        start = 0
        while start < len(words):
            end = min(start + target_words, len(words))
            chunks.append(PageText(page.page_number, " ".join(words[start:end])))
            if end == len(words):
                break
            start = max(end - overlap, start + 1)
    return chunks


def process_document(db: Session, document: KnowledgeDocument) -> KnowledgeDocument:
    try:
        pages = extract_pages(document.storage_key, document.content_type)
        chunks = chunk_pages(pages)
        if not chunks:
            raise InvalidDocument("No readable text was found in this document")
        for position, chunk in enumerate(chunks):
            db.add(
                DocumentChunk(
                    document_id=document.id,
                    organization_id=document.organization_id,
                    position=position,
                    page_number=chunk.page_number,
                    content=chunk.text,
                    token_count=len(chunk.text.split()),
                )
            )
        document.page_count = len(pages)
        document.chunk_count = len(chunks)
        document.status = DocumentStatus.ready
        db.commit()
        db.refresh(document)
        return document
    except Exception as exc:
        db.rollback()
        persisted = db.get(KnowledgeDocument, document.id)
        if persisted is not None:
            persisted.status = DocumentStatus.failed
            persisted.error_message = str(exc)[:500]
            db.commit()
            db.refresh(persisted)
            return persisted
        raise


def tokens(value: str) -> list[str]:
    return [token.lower() for token in TOKEN_RE.findall(value)]


def relevance_score(query: str, content: str) -> float:
    query_tokens = [token for token in tokens(query) if token not in QUERY_STOP_WORDS]
    content_tokens = tokens(content)
    if not query_tokens or not content_tokens:
        return 0.0
    frequencies = Counter(content_tokens)
    unique_query = set(query_tokens)
    matched = sum(min(frequencies[token], 3) for token in unique_query)
    coverage = sum(1 for token in unique_query if token in frequencies) / len(unique_query)
    density = matched / math.sqrt(len(content_tokens))
    phrase_bonus = 0.35 if query.lower() in content.lower() else 0.0
    return round((coverage * 0.7) + (min(density, 1.0) * 0.3) + phrase_bonus, 4)


def search_documents(
    db: Session,
    *,
    organization_id: uuid.UUID,
    query: str,
    limit: int,
    document_ids: list[uuid.UUID] | None = None,
) -> list[Citation]:
    statement = (
        select(DocumentChunk, KnowledgeDocument)
        .join(KnowledgeDocument, KnowledgeDocument.id == DocumentChunk.document_id)
        .where(
            DocumentChunk.organization_id == organization_id,
            KnowledgeDocument.status == DocumentStatus.ready,
        )
    )
    if document_ids:
        statement = statement.where(KnowledgeDocument.id.in_(document_ids))

    ranked: list[tuple[float, DocumentChunk, KnowledgeDocument]] = []
    for chunk, document in db.execute(statement):
        score = relevance_score(query, chunk.content)
        if score > 0:
            ranked.append((score, chunk, document))
    ranked.sort(key=lambda item: item[0], reverse=True)

    return [
        Citation(
            document_id=document.id,
            document_name=document.name,
            page_number=chunk.page_number,
            chunk_id=chunk.id,
            excerpt=chunk.content[:600],
            score=score,
        )
        for score, chunk, document in ranked[:limit]
    ]


def compose_grounded_answer(query: str, citations: list[Citation]) -> str:
    candidates: list[tuple[float, str]] = []
    seen: set[str] = set()
    for citation in citations[:3]:
        for sentence in re.split(r"(?<=[.!?])\s+", citation.excerpt):
            cleaned = sentence.strip()
            fingerprint = cleaned.lower()
            if len(cleaned) < 25 or fingerprint in seen:
                continue
            seen.add(fingerprint)
            candidates.append((relevance_score(query, cleaned), cleaned))
    candidates.sort(key=lambda item: item[0], reverse=True)
    selected = [sentence for score, sentence in candidates[:4] if score > 0]
    if not selected:
        return citations[0].excerpt.rstrip(" .") + "."
    return " ".join(selected)


def remove_stored_document(document: KnowledgeDocument) -> None:
    path = Path(document.storage_key)
    if path.exists():
        path.unlink()
    parent = path.parent
    if parent.exists() and not any(parent.iterdir()):
        shutil.rmtree(parent)
